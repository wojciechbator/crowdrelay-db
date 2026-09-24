from __future__ import annotations

import csv
import json
import os
import re
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import date
from pathlib import Path
from urllib.parse import parse_qs, quote_plus, unquote, urlparse

import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook

REPO_DB = Path("database.xlsx")
AUDIT_DIR = Path("audit")
TODAY = date.today().isoformat()
USER_AGENT = "CrowdRelayDB-ResearchAudit/3.0"

COUNTRY_FROM_SOURCE = {
    "/country/Poland": "Poland",
    "/country/Germany": "Germany",
    "/country/Czech-Republic": "Czechia",
    "/country/Slovakia": "Slovakia",
}

NON_BAND_TERMS = re.compile(
    r"(festival|festiwal|radio|turbo top|music city|\bartyst[a-z]*\b|artists?\b|"
    r"tour\s+20\d\d|\bclub\b|\bklub\b|\bpub\b|\bvenue\b|\bbooking\b|"
    r"\bpromotion\b|\bagency\b|music review|podcast|magazine|media)",
    re.I,
)

INACTIVE_TERMS = re.compile(
    r"\b(disbanded|split[- ]?up|dissolved|defunct|on hold|inactive|ceased|broke up)\b",
    re.I,
)

ACTIVE_TERMS = re.compile(
    r"\b(2026|2025|upcoming|tour|concert|show|festival|album|single|release|"
    r"live|gig|anniversary|present|active)\b",
    re.I,
)

GOOD_DOMAINS = {
    "metal-archives.com",
    "bandsintown.com",
    "songkick.com",
    "bandcamp.com",
    "facebook.com",
    "instagram.com",
    "youtube.com",
    "soundcloud.com",
    "reverbnation.com",
    "bandzone.cz",
    "last.fm",
    "genius.com",
}

BAD_DOMAINS = {
    "wikipedia.org",
    "dictionary.com",
    "merriam-webster.com",
    "walmart.com",
    "un.org",
    "tiktok.com",
    "websudoku.com",
}


def norm(v: object) -> str:
    return re.sub(r"\s+", " ", str(v or "").strip()).casefold()


def infer_country(source: str) -> str:
    for needle, country in COUNTRY_FROM_SOURCE.items():
        if needle in source:
            return country
    return ""


def decode_ddg_url(url: str) -> str:
    if not url:
        return ""
    p = urlparse(url)
    if "duckduckgo.com" not in p.netloc:
        return unquote(url)
    target = parse_qs(p.query).get("uddg", [""])[0]
    return unquote(target or url)


def result_domain(url: str) -> str:
    try:
        h = urlparse(url).netloc.casefold()
        return h.removeprefix("www.")
    except Exception:
        return ""


def relevant(name: str, item: dict, query_kind: str) -> tuple[bool, int]:
    title = norm(item.get("title"))
    snippet = norm(item.get("snippet"))
    url = item.get("url") or ""
    domain = result_domain(url)
    n = norm(name)

    # Word-boundary match, not substring: "Cult" must not KEEP a page titled
    # "Culture events in Berlin" — short names inside longer words were the
    # source of false-positive KEEPs.
    name_in_title = bool(re.search(r"\b" + re.escape(n) + r"\b", title)) if n else False
    if not (n == title or name_in_title or (len(n) >= 6 and n in snippet)):
        return False, 0

    if domain in BAD_DOMAINS or any(domain.endswith("." + d) for d in BAD_DOMAINS):
        return False, 0

    known_music_domain = any(
        domain == d or domain.endswith("." + d) for d in GOOD_DOMAINS
    )
    music_context = bool(re.search(
        r"\b(band|metal|metalcore|deathcore|hardcore|rock|punk|album|ep|single|"
        r"discography|concert|tour|live|gig|guitar|drums|vocal)\b",
        f"{title} {snippet}",
        re.I,
    ))
    if not known_music_domain and not music_context:
        return False, 0

    score = 0
    if n == title:
        score += 8
    elif n in title:
        score += 6
    else:
        score += 3
    if known_music_domain:
        score += 5
    if domain == "metal-archives.com" or domain.endswith(".metal-archives.com"):
        score += 8
    if "2026" in title or "2026" in snippet:
        score += 4
    if "2025" in title or "2025" in snippet:
        score += 2
    if music_context:
        score += 3

    return score >= 10, score


def ddg(session: requests.Session, query: str) -> list[dict]:
    url = "https://html.duckduckgo.com/html/?q=" + quote_plus(query)
    try:
        r = session.get(url, timeout=15)
    except requests.RequestException:
        return []
    if r.status_code != 200:
        return []

    soup = BeautifulSoup(r.text, "html.parser")
    out = []
    for item in soup.select(".result"):
        a = item.select_one("a.result__a")
        if not a:
            continue
        snippet_el = item.select_one(".result__snippet")
        real_url = decode_ddg_url(a.get("href", ""))
        if result_domain(real_url) in {"duckduckgo.com", ""}:
            continue
        out.append({
            "title": a.get_text(" ", strip=True),
            "url": real_url,
            "snippet": snippet_el.get_text(" ", strip=True) if snippet_el else "",
        })
    return out


def inspect(rec: dict) -> dict:
    session = requests.Session()
    session.headers.update({
        "User-Agent": USER_AGENT,
        "Accept-Language": "en-US,en;q=0.8",
    })
    name = str(rec.get("Name") or "").strip()
    source = str(rec.get("Source_URL") or "").strip()
    inferred_country = str(rec.get("Country") or "").strip() or infer_country(source)
    genre = str(rec.get("Genre") or "").strip()
    activity = str(rec.get("Activity") or "").strip()

    if not name:
        return {
            **rec,
            "Decision": "REVIEW",
            "Decision_Reason": "Empty name",
            "Verified_Source": source,
            "Verified_Date": TODAY,
            "Verified_Country": inferred_country,
            "Verified_Genre": genre,
            "Verified_Status": "",
            "Confidence_New": "Low",
            "Evidence": "empty name",
        }

    if NON_BAND_TERMS.search(name):
        return {
            **rec,
            "Decision": "REMOVE",
            "Decision_Reason": "Entity name indicates event/media/venue/organisation rather than a band",
            "Verified_Source": source,
            "Verified_Date": TODAY,
            "Verified_Country": inferred_country,
            "Verified_Genre": genre,
            "Verified_Status": "",
            "Confidence_New": "High",
            "Evidence": "entity-type rule",
        }

    src_lower = source.casefold()

    # Concrete current source already attached to the row.
    source_current = (
        "metalunderground.com" not in src_lower
        and bool(re.search(r"20(25|26)", source + " " + activity, re.I))
    )
    if source_current:
        return {
            **rec,
            "Decision": "KEEP",
            "Decision_Reason": "Existing row already has a concrete 2025/2026 source or activity marker",
            "Verified_Source": source,
            "Verified_Date": TODAY,
            "Verified_Country": inferred_country,
            "Verified_Genre": genre,
            "Verified_Status": "Active",
            "Confidence_New": "Medium",
            "Evidence": "existing current source/activity",
        }

    # One broad current query; this often returns Metal Archives, Bandsintown,
    # Bandcamp, Facebook, Instagram, YouTube and current event pages together.
    current_results = ddg(
        session,
        f'"{name}" 2026 band {inferred_country}'.strip(),
    )
    current_good = []
    for x in current_results:
        ok, score = relevant(name, x, "current")
        if ok:
            current_good.append((score, x))
    current_good.sort(reverse=True, key=lambda z: z[0])

    cur_blob = " ".join(f"{x['title']} {x['snippet']}" for _, x in current_good[:10])
    current_signal = bool(
        re.search(r"\b2026\b", cur_blob, re.I)
        and re.search(
            r"\b(concert|tour|show|festival|album|single|release|live|gig|band|metal|hardcore|metalcore)\b",
            cur_blob,
            re.I,
        )
    )

    ma_result = next(
        (
            x for _, x in current_good
            if result_domain(x["url"]) == "metal-archives.com"
            or result_domain(x["url"]).endswith(".metal-archives.com")
        ),
        None,
    )
    ma_title_active = bool(
        ma_result
        and re.search(r"\b(Active|active)\b", ma_result["snippet"] + " " + ma_result["title"])
    )

    # Only ask the status-focused query if the first search found no decisive current activity.
    if not current_signal and not ma_title_active:
        status_results = ddg(
            session,
            f'"{name}" site:metal-archives.com/bands "Status" "Active"',
        )
        status_good = []
        for x in status_results:
            ok, score = relevant(name, x, "ma")
            if ok:
                status_good.append((score, x))
        status_good.sort(reverse=True, key=lambda z: z[0])
        if status_good:
            top = status_good[0][1]
            snippet = top["snippet"] + " " + top["title"]
            if re.search(r"\b(active|present)\b", snippet, re.I):
                ma_result = top
                ma_title_active = True

    if current_signal and ma_title_active:
        confidence = "High"
    elif current_signal or ma_title_active:
        confidence = "Medium"
    else:
        confidence = "Low"

    if current_signal or ma_title_active:
        verified_source = ma_result["url"] if ma_result else current_good[0][1]["url"]
        return {
            **rec,
            "Decision": "KEEP",
            "Decision_Reason": "Band identity and current/active evidence found",
            "Verified_Source": verified_source,
            "Verified_Date": TODAY,
            "Verified_Country": inferred_country,
            "Verified_Genre": genre,
            "Verified_Status": "Active",
            "Confidence_New": confidence,
            "Evidence": (
                ("Metal Archives active; " if ma_title_active else "")
                + ("2026 current activity; " if current_signal else "")
            ).strip("; "),
        }

    # Never auto-delete a real-looking band based on absence of evidence.
    return {
        **rec,
        "Decision": "REVIEW",
        "Decision_Reason": "No sufficiently reliable current activity signal after exact-name/source filtering",
        "Verified_Source": source,
        "Verified_Date": TODAY,
        "Verified_Country": inferred_country,
        "Verified_Genre": genre,
        "Verified_Status": "",
        "Confidence_New": "Low",
        "Evidence": "no decisive current signal",
    }


def load_prior_audited_names() -> set[str]:
    audited: set[str] = set()
    search_roots = [Path("updates/applied"), Path("audit")]
    patterns = [
        "Audit_Peer_Bands__*.csv",
        "peer_bands_audit_batch__*.csv",
    ]

    for root in search_roots:
        if not root.exists():
            continue
        for pattern in patterns:
            for path in root.rglob(pattern):
                try:
                    with path.open("r", encoding="utf-8-sig", newline="") as fh:
                        for row in csv.reader(fh):
                            if row and row[0].strip() and row[0].strip().casefold() != "name":
                                audited.add(norm(row[0]))
                except OSError:
                    continue
    return audited


def update_batch_in_place(ws, headers: list[str], batch_results: list[dict]) -> tuple[int, int, int]:
    idx = {h: i + 1 for i, h in enumerate(headers)}
    name_col = idx["Name"]

    rows_by_name: dict[str, int] = {}
    for row_no in range(3, ws.max_row + 1):
        key = norm(ws.cell(row_no, name_col).value)
        if key:
            rows_by_name[key] = row_no

    kept = removed = reviewed = 0
    for r in batch_results:
        key = norm(r["Name"])
        row_no = rows_by_name.get(key)

        if r["Decision"] == "REMOVE":
            if row_no is not None:
                ws.delete_rows(row_no, 1)
                removed += 1
                # Rebuild row index after deletion because row numbers shift.
                rows_by_name = {
                    norm(ws.cell(n, name_col).value): n
                    for n in range(3, ws.max_row + 1)
                    if norm(ws.cell(n, name_col).value)
                }
            continue

        if row_no is None:
            # The source record should normally exist. If another job removed it,
            # preserve the audited data rather than silently dropping it.
            ws.append([r[h] if r[h] is not None else "" for h in headers])
            row_no = ws.max_row
            rows_by_name[key] = row_no

        values = [r[h] if r[h] is not None else "" for h in headers]
        values[idx["Country"] - 1] = r["Verified_Country"] or values[idx["Country"] - 1]
        values[idx["Genre"] - 1] = r["Verified_Genre"] or values[idx["Genre"] - 1]
        values[idx["Research_Date"] - 1] = TODAY
        # An inconclusive re-audit must not downgrade fields a previous audit
        # verified — only fill the marker when the cell is empty. REMOVE rows
        # are deleted separately; REVIEW keeps what was already known.
        review_only = r["Decision"] != "KEEP"
        values[idx["Status"] - 1] = r["Verified_Status"] or values[idx["Status"] - 1] or (
            "Needs verification" if review_only else "Active"
        )
        values[idx["Activity"] - 1] = (
            "2025/2026 activity verified"
            if r["Decision"] == "KEEP"
            else values[idx["Activity"] - 1] or "Needs current activity verification"
        )
        values[idx["Confidence"] - 1] = (
            r["Confidence_New"] if r["Decision"] == "KEEP"
            else values[idx["Confidence"] - 1] or "Low"
        )
        values[idx["Contact_Source"] - 1] = r["Verified_Source"] or values[idx["Contact_Source"] - 1]

        for col_no, value in enumerate(values, start=1):
            ws.cell(row_no, col_no).value = value

        if r["Decision"] == "KEEP":
            kept += 1
        else:
            reviewed += 1

    return kept, removed, reviewed


def main() -> None:
    wb = load_workbook(REPO_DB)
    ws = wb["Peer Bands"]
    headers = [c.value for c in ws[2]]
    records = [
        {h: row[i] if i < len(row) else "" for i, h in enumerate(headers)}
        for row in ws.iter_rows(min_row=3, values_only=True)
        if any(str(v or "").strip() for v in row)
    ]

    prior_audited = load_prior_audited_names()
    candidates = [r for r in records if norm(r.get("Name")) not in prior_audited]
    batch_size = int(os.environ.get("AUDIT_BATCH_SIZE", "10"))
    batch = candidates[:batch_size]

    print(
        f"PEER_AUDIT_START rows={len(records)} already_audited={len(prior_audited)} "
        f"remaining={len(candidates)} batch={len(batch)}"
    )

    if not batch:
        print("PEER_AUDIT_DONE no remaining candidates")
        return

    results = [None] * len(batch)
    with ThreadPoolExecutor(max_workers=min(batch_size, 10)) as ex:
        futures = {ex.submit(inspect, rec): i for i, rec in enumerate(batch)}
        done = 0
        for fut in as_completed(futures):
            i = futures[fut]
            try:
                results[i] = fut.result()
            except Exception as e:
                rec = batch[i]
                results[i] = {
                    **rec,
                    "Decision": "REVIEW",
                    "Decision_Reason": f"Audit error: {type(e).__name__}: {e}",
                    "Verified_Source": rec["Source_URL"],
                    "Verified_Date": TODAY,
                    "Verified_Country": rec["Country"],
                    "Verified_Genre": rec["Genre"],
                    "Verified_Status": "",
                    "Confidence_New": "Low",
                    "Evidence": "audit exception",
                }
            done += 1
            print(f"PEER_AUDIT_PROGRESS {done}/{len(batch)}")

    batch_id = len(prior_audited) + 1
    report = AUDIT_DIR / f"peer_bands_audit_batch__{batch_id:04d}.csv"
    out_headers = headers + [
        "Decision","Decision_Reason","Verified_Source","Verified_Date",
        "Verified_Country","Verified_Genre","Verified_Status","Confidence_New","Evidence"
    ]
    AUDIT_DIR.mkdir(exist_ok=True)
    with report.open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=out_headers)
        w.writeheader()
        w.writerows(results)

    summary = {
        "date": TODAY,
        "batch_id": batch_id,
        "rows": len(results),
        "already_audited_before_batch": len(prior_audited),
        "remaining_before_batch": len(candidates),
        "KEEP": sum(r["Decision"] == "KEEP" for r in results),
        "REMOVE": sum(r["Decision"] == "REMOVE" for r in results),
        "REVIEW": sum(r["Decision"] == "REVIEW" for r in results),
        "HIGH": sum(r["Confidence_New"] == "High" for r in results),
        "MEDIUM": sum(r["Confidence_New"] == "Medium" for r in results),
        "LOW": sum(r["Confidence_New"] == "Low" for r in results),
    }
    summary_path = AUDIT_DIR / f"peer_bands_audit_batch_summary__{batch_id:04d}.json"
    summary_path.write_text(
        json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print("PEER_AUDIT_SUMMARY", json.dumps(summary, ensure_ascii=False))

    if os.environ.get("APPLY_AUDIT") == "1":
        kept, removed, reviewed = update_batch_in_place(ws, headers, results)
        tmp = REPO_DB.with_name(REPO_DB.name + ".tmp")
        wb.save(tmp)
        os.replace(tmp, REPO_DB)
        print(
            f"PEER_AUDIT_APPLIED kept={kept} removed={removed} review={reviewed}"
        )

    print("PEER_AUDIT_DONE")


if __name__ == "__main__":
    main()
