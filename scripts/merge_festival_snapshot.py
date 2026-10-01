#!/usr/bin/env python3
from __future__ import annotations

import json
import os
import re
from pathlib import Path
from urllib.parse import urlparse

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
DB = ROOT / "database.xlsx"
PAYLOAD = Path(os.environ.get("SCOUT_PL_PAYLOAD", ROOT / "scout_runs" / "current.json"))

SHEET_MAP = {
    "RUN_INFO": "Festival Run Info",
    "OPPORTUNITIES": "Festival Opportunities",
    "ORGANIZERS": "Festival Organizers",
    "CONTACTS": "Festival Contacts",
    "CANONICAL_APPEND": "Festival Canonical Append",
    "SUPPORT_TARGETS": "Festival Support Targets",
    "OUTREACH_LOG": "Festival Outreach Log",
}

REQUIRED = {"RUN_INFO", "OPPORTUNITIES", "ORGANIZERS", "CONTACTS", "CANONICAL_APPEND"}


def norm(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip()).casefold()


def canon_url(value: object) -> str:
    raw = str(value or "").strip()
    if not raw:
        return ""
    try:
        p = urlparse(raw if "://" in raw else "//" + raw)
        host = p.netloc.casefold().removeprefix("www.")
        path = re.sub(r"/+$", "", p.path or "").casefold()
        return host + path
    except Exception:
        return norm(raw)


def canonical_contact_keys(wb) -> set[tuple[str, ...]]:
    if "Contacts" not in wb.sheetnames:
        return set()
    ws = wb["Contacts"]
    header_row = None
    headers = {}
    for r in range(1, min(ws.max_row, 10) + 1):
        vals = [norm(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)]
        if "email" in vals and "city" in vals:
            header_row = r
            headers = {v: i + 1 for i, v in enumerate(vals) if v}
            break
    if header_row is None:
        return set()

    out: set[tuple[str, ...]] = set()
    for r in range(header_row + 1, ws.max_row + 1):
        email = norm(ws.cell(r, headers.get("email", 1)).value)
        city = norm(ws.cell(r, headers.get("city", 1)).value)
        name = norm(ws.cell(r, headers.get("name", 1)).value)
        org = norm(ws.cell(r, headers.get("organization", 1)).value)
        if email and city:
            out.add(("email_city", email, city))
        if email:
            out.add(("email", email))
        if name and org:
            out.add(("name_org", name, org))
    return out


def row_identity(sheet: str, headers: list[str], row: list[object]) -> tuple[str, ...]:
    idx = {norm(h): i for i, h in enumerate(headers)}
    def get(name: str) -> str:
        i = idx.get(norm(name))
        return norm(row[i]) if i is not None and i < len(row) else ""
    def get_url(name: str) -> str:
        i = idx.get(norm(name))
        return canon_url(row[i]) if i is not None and i < len(row) else ""

    if sheet == "RUN_INFO":
        return ("run", get("Run_Date"), get("Scope"))
    if sheet == "OPPORTUNITIES":
        name = get("Name")
        organizer = get("Organizer")
        if name and organizer:
            return ("name_org", name, organizer)
        source = get_url("Source_URL")
        if source:
            return ("url", source)
        return ("name", name)
    if sheet == "ORGANIZERS":
        website = get_url("Website")
        email = get("Email")
        if website:
            return ("url", website)
        if email:
            return ("email", email)
        return ("org_city", get("Organization"), get("City"))
    if sheet == "CONTACTS":
        email = get("Email")
        city = get("City")
        if email and city:
            return ("email_city", email, city)
        if email:
            return ("email", email)
        return ("name_org_city", get("Name"), get("Organization/Event"), city)
    if sheet == "CANONICAL_APPEND":
        target = get("Target_Sheet")
        email = get("Email")
        city = get("City")
        if email and city:
            return ("target_email_city", target, email, city)
        if email:
            return ("target_email", target, email)
        source = get_url("Source_URL")
        if source:
            return ("target_url", target, source)
        return ("target_name_city", target, get("Name"), city)
    if sheet == "OUTREACH_LOG":
        email = get("Public Email")
        subject = get("Subject / Thread")
        recipient = get("Recipient / Organization")
        date = get("Date")
        if email and subject:
            return ("email_subject", email, subject)
        return ("recipient_subject_date", recipient, subject, date)
    if sheet == "SUPPORT_TARGETS":
        dedupe_key = get("Dedupe Key")
        if dedupe_key:
            return ("dedupe_key", dedupe_key)
        artist = get("Band / Artist")
        region = get("Region")
        if artist or region:
            return ("artist_region", artist, region)
        source = get_url("Public URL")
        if source:
            return ("url", source)
        return ("row",) + tuple(norm(x) for x in row)
    return tuple(norm(x) for x in row)


def dedupe_rows(sheet: str, headers: list[str], rows: list[list[object]]) -> list[list[object]]:
    deduped: dict[tuple[str, ...], list[object]] = {}
    order: list[tuple[str, ...]] = []
    for raw in rows:
        row = list(raw[: len(headers)]) + [""] * max(0, len(headers) - len(raw))
        if not any(norm(x) for x in row):
            continue
        key = row_identity(sheet, headers, row)
        if not any(key[1:]) and sheet != "RUN_INFO":
            continue
        if key not in deduped:
            deduped[key] = row
            order.append(key)
            continue
        current = deduped[key]
        for i, value in enumerate(row):
            if norm(value):
                current[i] = value
    return [deduped[k] for k in order]


def normalize_canonical_actions(headers: list[str], rows: list[list[object]], canonical_contacts: set[tuple[str, ...]]) -> list[list[object]]:
    idx = {norm(h): i for i, h in enumerate(headers)}
    action_i = idx.get("action")
    target_i = idx.get("target_sheet")
    email_i = idx.get("email")
    city_i = idx.get("city")
    if None in (action_i, target_i, email_i, city_i):
        return rows

    out = []
    for row in rows:
        row = list(row)
        if norm(row[target_i]) == "contacts":
            email = norm(row[email_i])
            city = norm(row[city_i])
            exists = (email and city and ("email_city", email, city) in canonical_contacts) or (
                email and ("email", email) in canonical_contacts
            )
            if exists and norm(row[action_i]) == "new_pending":
                row[action_i] = "UPDATE_PENDING"
        out.append(row)
    return out


def find_header(ws) -> tuple[int, dict[str, int]]:
    for r in range(1, min(ws.max_row, 10) + 1):
        vals = [norm(ws.cell(r, col).value) for col in range(1, ws.max_column + 1)]
        if "name" in vals or "email" in vals:
            return r, {v: i + 1 for i, v in enumerate(vals) if v}
    return 1, {}


def canonical_row_key(sheet: str, ws, r: int, idx: dict[str, int]) -> tuple[str, ...] | None:
    def v(name: str) -> str:
        col = idx.get(name)
        return norm(ws.cell(r, col).value) if col else ""
    def u(name: str) -> str:
        col = idx.get(name)
        return canon_url(ws.cell(r, col).value) if col else ""

    if sheet == "Venues":
        name, city = v("name"), v("city")
        return ("name_city", name, city) if name and city else None
    if sheet == "Peer Bands":
        name, city = v("name"), v("city")
        return ("name_city", name, city) if name else None
    if sheet == "Beacons":
        kind, city, dest = v("kind"), v("city"), u("destination_url")
        return ("kind_city_url", kind, city, dest) if kind and city and dest else None
    if sheet == "Booking Agents":
        name, agency = v("name"), v("agency")
        if name and agency:
            return ("name_agency", name, agency)
        return ("name", name) if name else None
    if sheet == "Contacts":
        email, city, name, org = v("email"), v("city"), v("name"), v("organization")
        if email and city:
            return ("email_city", email, city)
        if email:
            return ("email", email)
        if name and org:
            return ("name_org", name, org)
    return None


def dedupe_canonical_workbook(wb) -> dict[str, int]:
    removed: dict[str, int] = {}
    for sheet in ("Venues", "Peer Bands", "Beacons", "Booking Agents", "Contacts"):
        if sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]
        header_row, idx = find_header(ws)
        seen: dict[tuple[str, ...], int] = {}
        duplicates: list[tuple[int, int]] = []
        for r in range(header_row + 1, ws.max_row + 1):
            key = canonical_row_key(sheet, ws, r, idx)
            if key is None:
                continue
            first = seen.get(key)
            if first is None:
                seen[key] = r
            else:
                duplicates.append((first, r))

        if not duplicates:
            continue

        for first, dup in duplicates:
            for col in range(1, ws.max_column + 1):
                if not norm(ws.cell(first, col).value) and norm(ws.cell(dup, col).value):
                    ws.cell(first, col).value = ws.cell(dup, col).value

        for _, dup in sorted(duplicates, key=lambda p: p[1], reverse=True):
            ws.delete_rows(dup, 1)
        removed[sheet] = len(duplicates)
    return removed


def read_existing_rows(wb, title: str, headers: list[str]) -> list[list[object]]:
    if title not in wb.sheetnames:
        return []
    ws = wb[title]
    actual_headers = [str(ws.cell(1, c).value or "") for c in range(1, len(headers) + 1)]
    if actual_headers != headers:
        raise SystemExit(
            f"{title}: existing header drift; expected {headers!r}, got {actual_headers!r}"
        )
    rows: list[list[object]] = []
    for r in range(2, ws.max_row + 1):
        row = [ws.cell(r, c).value for c in range(1, len(headers) + 1)]
        if any(norm(x) for x in row):
            rows.append(row)
    return rows


def merge_state_rows(
    sheet: str,
    headers: list[str],
    existing_rows: list[list[object]],
    incoming_rows: list[list[object]],
) -> list[list[object]]:
    existing = dedupe_rows(sheet, headers, existing_rows)
    incoming = dedupe_rows(sheet, headers, incoming_rows)

    merged: dict[tuple[str, ...], list[object]] = {}
    order: list[tuple[str, ...]] = []

    for row in existing:
        key = row_identity(sheet, headers, row)
        if key not in merged:
            merged[key] = list(row)
            order.append(key)

    for row in incoming:
        key = row_identity(sheet, headers, row)
        if key not in merged:
            merged[key] = list(row)
            order.append(key)
            continue
        current = merged[key]
        for i, value in enumerate(row):
            if norm(value):
                current[i] = value

    return [merged[k] for k in order]


def replace_sheet(wb, title: str, headers: list[str], rows: list[list[object]]) -> None:
    if title in wb.sheetnames:
        old = wb[title]
        idx = wb.sheetnames.index(title)
        wb.remove(old)
        ws = wb.create_sheet(title, idx)
    else:
        ws = wb.create_sheet(title)

    ws.append(headers)
    for row in rows:
        ws.append(row)

    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col in range(1, ws.max_column + 1):
        max_len = max((len(str(ws.cell(r, col).value or "")) for r in range(1, min(ws.max_row, 200) + 1)), default=8)
        ws.column_dimensions[get_column_letter(col)].width = min(max(max_len + 2, 10), 48)


def main() -> None:
    if not DB.exists():
        raise SystemExit("database.xlsx missing")
    payload = json.loads(PAYLOAD.read_text(encoding="utf-8"))
    if payload.get("snapshot_mode") != "current_run_only":
        raise SystemExit("snapshot_mode must be current_run_only")
    sheets = payload.get("sheets") or {}
    missing = REQUIRED - set(sheets)
    if missing:
        raise SystemExit(f"missing payload sheets: {sorted(missing)}")

    wb = load_workbook(DB)
    canonical_deduped = dedupe_canonical_workbook(wb)
    canonical_contacts = canonical_contact_keys(wb)
    stats = {}

    for source_name, target_name in SHEET_MAP.items():
        if source_name not in sheets:
            continue
        source = sheets[source_name]
        headers = list(source.get("headers") or [])
        incoming_rows = [list(r) for r in (source.get("rows") or [])]
        existing_rows = read_existing_rows(wb, target_name, headers)
        rows = merge_state_rows(source_name, headers, existing_rows, incoming_rows)
        if source_name == "CANONICAL_APPEND":
            rows = normalize_canonical_actions(headers, rows, canonical_contacts)
        replace_sheet(wb, target_name, headers, rows)
        stats[target_name] = {
            "before": len(dedupe_rows(source_name, headers, existing_rows)),
            "incoming": len(dedupe_rows(source_name, headers, incoming_rows)),
            "after": len(rows),
        }

    tmp = DB.with_suffix(".tmp.xlsx")
    wb.save(tmp)
    os.replace(tmp, DB)
    print("FESTIVAL_MERGE_OK", {"festival_rows": stats, "canonical_duplicates_removed": canonical_deduped})


if __name__ == "__main__":
    main()
