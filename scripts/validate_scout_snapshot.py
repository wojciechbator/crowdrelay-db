#!/usr/bin/env python3
from __future__ import annotations

import json
import re
from pathlib import Path
from urllib.parse import urlparse

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
DB = ROOT / "database.xlsx"
PAYLOAD = ROOT / "scout_runs" / "current.json"

SHEET_MAP = {
    "RUN_INFO": "Festival Run Info",
    "OPPORTUNITIES": "Festival Opportunities",
    "ORGANIZERS": "Festival Organizers",
    "CONTACTS": "Festival Contacts",
    "CANONICAL_APPEND": "Festival Canonical Append",
    "SUPPORT_TARGETS": "Festival Support Targets",
    "OUTREACH_LOG": "Festival Outreach Log",
}

FESTIVAL_SHEETS = {
    "Festival Run Info",
    "Festival Opportunities",
    "Festival Organizers",
    "Festival Contacts",
    "Festival Canonical Append",
}
OPTIONAL_FESTIVAL_SHEETS = {"Festival Support Targets", "Festival Outreach Log"}


def norm(v: object) -> str:
    return re.sub(r"\s+", " ", str(v or "").strip()).casefold()


def canon_url(v: object) -> str:
    raw = str(v or "").strip()
    if not raw:
        return ""
    p = urlparse(raw if "://" in raw else "//" + raw)
    return p.netloc.casefold().removeprefix("www.") + re.sub(r"/+$", "", p.path or "").casefold()


def fail(msg: str) -> None:
    raise SystemExit("UNIFIED_DATABASE_QA_FAIL: " + msg)


def header(ws) -> tuple[int, dict[str, int]]:
    for r in range(1, min(ws.max_row, 10) + 1):
        vals = [norm(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)]
        if any(vals):
            return r, {v: i + 1 for i, v in enumerate(vals) if v}
    fail(f"{ws.title}: no header")
    raise AssertionError


def canonical_key(sheet: str, ws, r: int, idx: dict[str, int]):
    def v(name: str) -> str:
        col = idx.get(name)
        return norm(ws.cell(r, col).value) if col else ""
    def u(name: str) -> str:
        col = idx.get(name)
        return canon_url(ws.cell(r, col).value) if col else ""

    if sheet == "Venues":
        name, city = v("name"), v("city")
        return ("name_city", name, city) if name and city else None
    if sheet == "Peer Bands":
        name, city = v("name"), v("city")
        return ("name_city", name, city) if name else None
    if sheet == "Beacons":
        kind, city, dest = v("kind"), v("city"), u("destination_url")
        return ("kind_city_url", kind, city, dest) if kind and city and dest else None
    if sheet == "Booking Agents":
        name, agency = v("name"), v("agency")
        if name and agency:
            return ("name_agency", name, agency)
        return ("name", name) if name else None
    if sheet == "Contacts":
        email, city, name, org = v("email"), v("city"), v("name"), v("organization")
        if email and city:
            return ("email_city", email, city)
        if email:
            return ("email", email)
        if name and org:
            return ("name_org", name, org)
    return None


def assert_unique_canonical(wb) -> None:
    for sheet in ("Venues", "Peer Bands", "Beacons", "Booking Agents", "Contacts"):
        ws = wb[sheet]
        hr, idx = header(ws)
        seen = {}
        for r in range(hr + 1, ws.max_row + 1):
            key = canonical_key(sheet, ws, r, idx)
            if key is None:
                continue
            if key in seen:
                fail(f"{sheet}: duplicate logical key rows {seen[key]} and {r}: {key}")
            seen[key] = r


def assert_unique_festival(wb) -> None:
    def rows(ws):
        hr, idx = header(ws)
        return hr, idx, range(hr + 1, ws.max_row + 1)

    ws = wb["Festival Opportunities"]; _, idx, rr = rows(ws)
    seen = set()
    for r in rr:
        source = canon_url(ws.cell(r, idx.get("source_url", 0)).value) if idx.get("source_url") else ""
        name = norm(ws.cell(r, idx.get("name", 0)).value) if idx.get("name") else ""
        org = norm(ws.cell(r, idx.get("organizer", 0)).value) if idx.get("organizer") else ""
        key = ("name_org", name, org) if name and org else (("url", source) if source else ("name", name))
        if key in seen:
            fail(f"Festival Opportunities duplicate: {key}")
        seen.add(key)

    ws = wb["Festival Organizers"]; _, idx, rr = rows(ws)
    seen = set()
    for r in rr:
        website = canon_url(ws.cell(r, idx.get("website", 0)).value) if idx.get("website") else ""
        email = norm(ws.cell(r, idx.get("email", 0)).value) if idx.get("email") else ""
        org = norm(ws.cell(r, idx.get("organization", 0)).value) if idx.get("organization") else ""
        city = norm(ws.cell(r, idx.get("city", 0)).value) if idx.get("city") else ""
        key = ("url", website) if website else (("email", email) if email else ("org_city", org, city))
        if key in seen:
            fail(f"Festival Organizers duplicate: {key}")
        seen.add(key)

    ws = wb["Festival Contacts"]; _, idx, rr = rows(ws)
    seen = set()
    for r in rr:
        email = norm(ws.cell(r, idx.get("email", 0)).value) if idx.get("email") else ""
        city = norm(ws.cell(r, idx.get("city", 0)).value) if idx.get("city") else ""
        name = norm(ws.cell(r, idx.get("name", 0)).value) if idx.get("name") else ""
        org = norm(ws.cell(r, idx.get("organization/event", 0)).value) if idx.get("organization/event") else ""
        key = ("email_city", email, city) if email and city else (("email", email) if email else ("name_org_city", name, org, city))
        if key in seen:
            fail(f"Festival Contacts duplicate: {key}")
        seen.add(key)


    if "Festival Outreach Log" in wb.sheetnames:
        ws = wb["Festival Outreach Log"]; _, idx, rr = rows(ws)
        seen = set()
        for r in rr:
            email = norm(ws.cell(r, idx.get("public email", 0)).value) if idx.get("public email") else ""
            subject = norm(ws.cell(r, idx.get("subject / thread", 0)).value) if idx.get("subject / thread") else ""
            recipient = norm(ws.cell(r, idx.get("recipient / organization", 0)).value) if idx.get("recipient / organization") else ""
            date = norm(ws.cell(r, idx.get("date", 0)).value) if idx.get("date") else ""
            key = ("email_subject", email, subject) if email and subject else ("recipient_subject_date", recipient, subject, date)
            if key in seen:
                fail(f"Festival Outreach Log duplicate: {key}")
            seen.add(key)


def assert_no_false_new_contacts(wb) -> None:
    ws = wb["Contacts"]; hr, idx = header(ws)
    keys = set()
    for r in range(hr + 1, ws.max_row + 1):
        email = norm(ws.cell(r, idx.get("email", 0)).value) if idx.get("email") else ""
        city = norm(ws.cell(r, idx.get("city", 0)).value) if idx.get("city") else ""
        if email and city:
            keys.add(("email_city", email, city))
        if email:
            keys.add(("email", email))

    ws = wb["Festival Canonical Append"]; hr, idx = header(ws)
    for r in range(hr + 1, ws.max_row + 1):
        target = norm(ws.cell(r, idx.get("target_sheet", 0)).value) if idx.get("target_sheet") else ""
        action = norm(ws.cell(r, idx.get("action", 0)).value) if idx.get("action") else ""
        email = norm(ws.cell(r, idx.get("email", 0)).value) if idx.get("email") else ""
        city = norm(ws.cell(r, idx.get("city", 0)).value) if idx.get("city") else ""
        exists = (email and city and ("email_city", email, city) in keys) or (email and ("email", email) in keys)
        if target == "contacts" and action == "new_pending" and exists:
            fail(f"Festival Canonical Append row {r}: existing contact incorrectly marked NEW_PENDING: {email}")


def payload_key(sheet: str, headers: list[str], row: list[object]) -> tuple[str, ...]:
    idx = {norm(h): i for i, h in enumerate(headers)}

    def get(name: str) -> str:
        i = idx.get(norm(name))
        return norm(row[i]) if i is not None and i < len(row) else ""

    def get_url(name: str) -> str:
        i = idx.get(norm(name))
        return canon_url(row[i]) if i is not None and i < len(row) else ""

    if sheet == "RUN_INFO":
        return ("run", get("Run_Date"), get("Scope"))
    if sheet == "OPPORTUNITIES":
        name, org = get("Name"), get("Organizer")
        if name and org:
            return ("name_org", name, org)
        source = get_url("Source_URL")
        return ("url", source) if source else ("name", name)
    if sheet == "ORGANIZERS":
        website, email = get_url("Website"), get("Email")
        if website:
            return ("url", website)
        if email:
            return ("email", email)
        return ("org_city", get("Organization"), get("City"))
    if sheet == "CONTACTS":
        email, city = get("Email"), get("City")
        if email and city:
            return ("email_city", email, city)
        if email:
            return ("email", email)
        return ("name_org_city", get("Name"), get("Organization/Event"), city)
    if sheet == "CANONICAL_APPEND":
        target, email, city = get("Target_Sheet"), get("Email"), get("City")
        if email and city:
            return ("target_email_city", target, email, city)
        if email:
            return ("target_email", target, email)
        source = get_url("Source_URL")
        return ("target_url", target, source) if source else ("target_name_city", target, get("Name"), city)
    if sheet == "OUTREACH_LOG":
        email, subject = get("Public Email"), get("Subject / Thread")
        recipient, date = get("Recipient / Organization"), get("Date")
        if email and subject:
            return ("email_subject", email, subject)
        return ("recipient_subject_date", recipient, subject, date)
    if sheet == "SUPPORT_TARGETS":
        dedupe_key = get("Dedupe Key")
        if dedupe_key:
            return ("dedupe_key", dedupe_key)
        artist, region = get("Band / Artist"), get("Region")
        if artist or region:
            return ("artist_region", artist, region)
        source = get_url("Public URL")
        return ("url", source) if source else ("row",) + tuple(norm(x) for x in row)
    return tuple(norm(x) for x in row)


def assert_north_star_payload_contract() -> None:
    payload = json.loads(PAYLOAD.read_text(encoding="utf-8"))
    sheets = payload.get("sheets") or {}

    required_headers = {
        "OPPORTUNITIES": {
            "Name", "Organizer", "Type", "Country", "City", "Current_Status",
            "Source_URL", "Submission_Method", "Relevance_Score", "Priority",
            "Verification", "Dedupe_Key", "Source_Checked", "Why_Fit", "Red_Flags",
        },
        "ORGANIZERS": {
            "Organization", "Type", "Website", "Email", "Verification",
            "Country", "Relevance_Score", "Dedupe_Key", "Source_Checked",
        },
    }

    for sheet_name, required in required_headers.items():
        source = sheets.get(sheet_name)
        if source is None:
            fail(f"{sheet_name}: missing from current-run payload")
        headers = list(source.get("headers") or [])
        header_set = set(headers)
        missing = required - header_set
        if missing:
            fail(f"{sheet_name}: north-star columns missing: {sorted(missing)}")

        idx = {h: i for i, h in enumerate(headers)}
        for row_number, raw in enumerate(source.get("rows") or [], start=2):
            row = list(raw) + [""] * max(0, len(headers) - len(raw))
            if not any(norm(v) for v in row):
                continue

            def value(name: str) -> str:
                return str(row[idx[name]] or "").strip()

            if sheet_name == "OPPORTUNITIES":
                for name in (
                    "Name", "Organizer", "Type", "Country", "City",
                    "Current_Status", "Source_URL", "Submission_Method",
                    "Priority", "Verification", "Dedupe_Key", "Source_Checked",
                    "Why_Fit",
                ):
                    if not value(name):
                        fail(f"{sheet_name} row {row_number}: required {name} is blank")
                try:
                    score = float(value("Relevance_Score"))
                except ValueError:
                    fail(f"{sheet_name} row {row_number}: Relevance_Score is not numeric")
                if not 0 <= score <= 100:
                    fail(f"{sheet_name} row {row_number}: Relevance_Score outside 0..100")
                if value("Priority").upper() not in {"A", "B", "C", "D"}:
                    fail(f"{sheet_name} row {row_number}: invalid Priority")
            else:
                for name in (
                    "Organization", "Type", "Country", "Verification",
                    "Dedupe_Key", "Source_Checked",
                ):
                    if not value(name):
                        fail(f"{sheet_name} row {row_number}: required {name} is blank")
                if not value("Website") and not value("Email"):
                    fail(f"{sheet_name} row {row_number}: organizer has no public route")
                try:
                    score = float(value("Relevance_Score"))
                except ValueError:
                    fail(f"{sheet_name} row {row_number}: Relevance_Score is not numeric")
                if not 0 <= score <= 100:
                    fail(f"{sheet_name} row {row_number}: Relevance_Score outside 0..100")


def assert_payload_landed(wb) -> None:
    if not PAYLOAD.exists():
        fail("scout_runs/current.json missing")
    payload = json.loads(PAYLOAD.read_text(encoding="utf-8"))
    if payload.get("snapshot_mode") != "current_run_only":
        fail("snapshot_mode must be current_run_only")

    for source_name, target_name in SHEET_MAP.items():
        source = (payload.get("sheets") or {}).get(source_name)
        if source is None:
            continue
        headers = list(source.get("headers") or [])
        incoming = [list(r) for r in (source.get("rows") or [])]
        if target_name not in wb.sheetnames:
            fail(f"{target_name}: missing after merge")
        ws = wb[target_name]
        actual_headers = [str(ws.cell(1, c).value or "") for c in range(1, len(headers) + 1)]
        if actual_headers != headers:
            fail(f"{target_name}: header drift after merge")

        workbook_rows = [
            [ws.cell(r, c).value for c in range(1, len(headers) + 1)]
            for r in range(2, ws.max_row + 1)
            if any(norm(ws.cell(r, c).value) for c in range(1, len(headers) + 1))
        ]
        by_key = {payload_key(source_name, headers, row): row for row in workbook_rows}

        for incoming_row in incoming:
            row = list(incoming_row[: len(headers)]) + [""] * max(0, len(headers) - len(incoming_row))
            if not any(norm(x) for x in row):
                continue
            key = payload_key(source_name, headers, row)
            landed = by_key.get(key)
            if landed is None:
                fail(f"{target_name}: payload row missing after merge: {key}")
            if source_name == "CANONICAL_APPEND":
                continue
            for i, value in enumerate(row):
                if norm(value) and norm(landed[i]) != norm(value):
                    fail(
                        f"{target_name}: payload value mismatch for {key}, "
                        f"column {headers[i]!r}: expected {value!r}, got {landed[i]!r}"
                    )


def main() -> None:
    if not DB.exists() or DB.stat().st_size == 0:
        fail("database.xlsx missing or empty")
    wb = load_workbook(DB, read_only=False)
    missing = FESTIVAL_SHEETS - set(wb.sheetnames)
    if missing:
        fail(f"missing merged festival sheets: {sorted(missing)}")
    assert_unique_canonical(wb)
    assert_unique_festival(wb)
    assert_no_false_new_contacts(wb)
    assert_north_star_payload_contract()
    assert_payload_landed(wb)
    print("UNIFIED_DATABASE_QA_OK", {
        "festival_sheets": sorted(FESTIVAL_SHEETS & set(wb.sheetnames)),
        "optional": sorted(OPTIONAL_FESTIVAL_SHEETS & set(wb.sheetnames)),
    })


if __name__ == "__main__":
    main()
